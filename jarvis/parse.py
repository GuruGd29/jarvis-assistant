"""Turn a file into sections: (heading, page, text). Pages are 1-based for PDFs, None otherwise."""
import html
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Section:
    heading: str | None
    page: int | None
    text: str


def parse_file(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in (".md", ".markdown"):
        return _parse_markdown(path.read_text(errors="replace"))
    if suffix in (".html", ".htm"):
        return [Section(None, None, _strip_html(path.read_text(errors="replace")))]
    return [Section(None, None, path.read_text(errors="replace"))]


def _parse_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(Section(None, number, text))
    return sections


def _parse_docx(path: Path) -> list[Section]:
    import docx

    document = docx.Document(str(path))
    sections: list[Section] = []
    heading, lines = None, []
    for para in document.paragraphs:
        if para.style.name.lower().startswith("heading") and para.text.strip():
            if lines:
                sections.append(Section(heading, None, "\n".join(lines)))
            heading, lines = para.text.strip(), []
        elif para.text.strip():
            lines.append(para.text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    if lines:
        sections.append(Section(heading, None, "\n".join(lines)))
    return sections


def _parse_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    heading, lines = None, []
    for line in text.splitlines():
        match = re.match(r"^#{1,6}\s+(.*)", line)
        if match:
            if "".join(lines).strip():
                sections.append(Section(heading, None, "\n".join(lines)))
            heading, lines = match.group(1).strip(), []
        else:
            lines.append(line)
    if "".join(lines).strip():
        sections.append(Section(heading, None, "\n".join(lines)))
    return sections


def _strip_html(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    return html.unescape(re.sub(r"<[^>]+>", " ", raw))
